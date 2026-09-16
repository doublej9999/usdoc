import asyncio
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
import json
import os
from pathlib import Path
import re
import sys
from threading import Lock
import time
from typing import Any, Dict, List
import uuid
import zipfile

from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from pydantic import BaseModel
from starlette.concurrency import run_in_threadpool

from config import (
    APP_TITLE,
    BASE_DIR,
    GENERATION_RECORDS_FILE,
    OUTPUT_DIR,
    STATIC_DIR,
    TEMPLATES_DIR,
    UPLOAD_DIR,
    ensure_runtime_directories,
)
from services.document_service import (
    convert_docx_to_pdf,
    generate_default_document,
    generate_document,
    validate_template,
)
from utils.excel import build_row_prompt, resolve_column_index
from utils.files import sanitize_filename, sanitize_output_filename

ensure_runtime_directories()

app = FastAPI(title=APP_TITLE)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
templates = Jinja2Templates(directory=str(TEMPLATES_DIR))


class GenerateRequest(BaseModel):
    prompt: str
    api_url: str
    api_key: str
    model: str
    template_name: str
    output_filename: str = ""


generation_records: Dict[str, Dict[str, Any]] = {}
generation_records_lock = Lock()


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def dump_generation_records() -> None:
    with generation_records_lock:
        records = list(generation_records.values())
        records.sort(key=lambda x: x.get("created_at", ""), reverse=True)
        payload = json.dumps(records, ensure_ascii=False, indent=2)

        # Write-then-rename: a crash mid-write must not truncate the whole history.
        temp_file = GENERATION_RECORDS_FILE.with_name(GENERATION_RECORDS_FILE.name + ".tmp")
        temp_file.write_text(payload, encoding="utf-8")
        os.replace(temp_file, GENERATION_RECORDS_FILE)


def load_generation_records() -> None:
    if not GENERATION_RECORDS_FILE.exists():
        return

    try:
        records = json.loads(GENERATION_RECORDS_FILE.read_text(encoding="utf-8"))
        if not isinstance(records, list):
            raise ValueError("记录文件格式不正确，期望 JSON 数组")
    except Exception as exc:
        print(f"[usdoc] 读取生成记录失败，已跳过: {exc}", file=sys.stderr)
        return

    interrupted = False
    current_time = now_iso()

    with generation_records_lock:
        generation_records.clear()
        for record in records:
            if not isinstance(record, dict) or not record.get("id"):
                continue
            # Background jobs live in this process only, so anything still
            # "pending"/"processing" died with the previous run.
            if record.get("status") in ("pending", "processing"):
                record["status"] = "failed"
                record["message"] = "Word 生成失败"
                record["error"] = "服务重启导致任务中断，请重新生成"
                record["updated_at"] = current_time
                interrupted = True
            generation_records[record["id"]] = record

    if interrupted:
        dump_generation_records()


def add_generation_record(req: GenerateRequest) -> Dict[str, Any]:
    record_id = uuid.uuid4().hex
    current_time = now_iso()
    record = {
        "id": record_id,
        "prompt": req.prompt,
        "api_url": req.api_url,
        "model": req.model,
        "template_name": req.template_name,
        "generated_filename": req.output_filename,
        "status": "pending",
        "message": "任务已创建，等待后台生成",
        "filename": "",
        "download_url": "",
        "error": "",
        "created_at": current_time,
        "updated_at": current_time,
    }

    with generation_records_lock:
        generation_records[record_id] = record

    dump_generation_records()
    return record


def update_generation_record(record_id: str, **fields) -> None:
    updated = False
    with generation_records_lock:
        record = generation_records.get(record_id)
        if record:
            record.update(fields)
            record["updated_at"] = now_iso()
            updated = True

    if updated:
        dump_generation_records()


def get_generation_record(record_id: str) -> Dict[str, Any]:
    with generation_records_lock:
        record = generation_records.get(record_id)
        if not record:
            raise HTTPException(status_code=404, detail="生成记录不存在")
        return dict(record)


def list_generation_records(limit: int) -> List[Dict[str, Any]]:
    with generation_records_lock:
        records = [dict(item) for item in generation_records.values()]

    records.sort(key=lambda x: x.get("created_at", ""), reverse=True)
    return records[:limit]


def generate_common(req: GenerateRequest) -> Dict[str, Any]:
    return generate_document(
        prompt=req.prompt,
        api_url=req.api_url,
        api_key=req.api_key,
        model=req.model,
        template_name=req.template_name,
        output_filename=req.output_filename,
        upload_dir=UPLOAD_DIR,
        output_dir=OUTPUT_DIR,
    )


def generate_default_docx(template_name: str, output_filename: str) -> str:
    return generate_default_document(
        template_name=template_name,
        output_filename=output_filename,
        upload_dir=UPLOAD_DIR,
        output_dir=OUTPUT_DIR,
    )


load_generation_records()


@app.get("/")
async def index(request: Request):
    return templates.TemplateResponse(request, "index.html", {"request": request})


@app.post("/upload-template")
async def upload_template(file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=400, detail="仅支持 .docx 模板")

    filename = sanitize_filename(file.filename)
    if not filename:
        raise HTTPException(status_code=400, detail="无效的文件名")
    save_path = UPLOAD_DIR / filename

    content = await file.read()
    with open(save_path, "wb") as f:
        f.write(content)

    return JSONResponse(
        {
            "message": "模板上传成功",
            "template_name": filename,
        }
    )


@app.get("/templates")
async def list_templates():
    templates_map: Dict[str, Dict[str, Any]] = {}

    if UPLOAD_DIR.exists():
        for f in sorted(UPLOAD_DIR.glob("*.docx")):
            if f.is_file() and not f.name.startswith("~$") and not f.name.startswith("."):
                stat = f.stat()
                templates_map[f.name] = {
                    "name": f.name,
                    "size": stat.st_size,
                    "updated_at": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
                }

    if BASE_DIR.exists():
        for f in sorted(BASE_DIR.glob("*.docx")):
            if (
                f.is_file()
                and not f.name.startswith("~$")
                and not f.name.startswith(".")
                and not f.name.endswith("_tmp_read.docx")
            ):
                if f.name not in templates_map:
                    stat = f.stat()
                    templates_map[f.name] = {
                        "name": f.name,
                        "size": stat.st_size,
                        "updated_at": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
                    }

    return {"templates": list(templates_map.values())}


@app.get("/templates/{template_name}/variables")
async def get_template_variables(template_name: str):
    target_name = template_name
    if not target_name.lower().endswith(".docx"):
        target_name = f"{target_name}.docx"

    try:
        _, variables = validate_template(target_name, UPLOAD_DIR)
    except HTTPException:
        if target_name != template_name:
            _, variables = validate_template(template_name, UPLOAD_DIR)
        else:
            raise

    return {"template_name": template_name, "variables": variables}


def run_docx_generation_task(record_id: str, request_payload: Dict[str, Any]) -> None:
    update_generation_record(
        record_id,
        status="processing",
        message="后台正在生成 Word",
        error="",
    )

    try:
        req = GenerateRequest(**request_payload)
        result = generate_common(req)
        filename = result["docx_filename"]
        update_generation_record(
            record_id,
            status="completed",
            message="Word 生成完成",
            generated_filename=filename,
            filename=filename,
            download_url=f"/download/{filename}",
            error="",
        )
    except HTTPException as exc:
        detail = exc.detail if isinstance(exc.detail, str) else json.dumps(exc.detail, ensure_ascii=False)
        update_generation_record(
            record_id,
            status="failed",
            message="Word 生成失败",
            error=detail,
        )
    except Exception as exc:
        update_generation_record(
            record_id,
            status="failed",
            message="Word 生成失败",
            error=str(exc),
        )


@app.post("/generate-docx")
async def generate_docx(req: GenerateRequest, background_tasks: BackgroundTasks):
    record = add_generation_record(req)
    payload = req.model_dump() if hasattr(req, "model_dump") else req.dict()
    background_tasks.add_task(run_docx_generation_task, record["id"], payload)

    return {
        "message": "Word 已提交后台生成",
        "record": record,
    }


def _process_excel_row(
    task_info: Dict[str, Any],
    api_url: str,
    api_key: str,
    model: str,
    template_name: str,
) -> Dict[str, Any]:
    row_num = task_info["row_num"]
    row_prompt = task_info["prompt"]
    output_name = task_info["output_name"]

    req = GenerateRequest(
        prompt=row_prompt,
        api_url=api_url,
        api_key=api_key,
        model=model,
        template_name=template_name,
        output_filename=output_name,
    )

    last_error = ""
    attempts = 0
    max_retries = 3

    for attempt in range(1, max_retries + 1):
        attempts = attempt
        try:
            result = generate_common(req)
            return {
                "row_num": row_num,
                "output_name": output_name,
                "filename": result["docx_filename"],
                "success": True,
                "fallback": False,
                "error": None,
                "attempts": attempts,
            }
        except Exception as exc:
            if isinstance(exc, HTTPException):
                detail = exc.detail if isinstance(exc.detail, str) else str(exc.detail)
                last_error = f"HTTP {exc.status_code}: {detail}"
                # Do NOT retry on 400/401/403/404 errors
                if exc.status_code in (400, 401, 403, 404):
                    break
            else:
                last_error = str(exc)

            if attempt < max_retries:
                time.sleep(0.5 * attempt)

    # Fallback to default document
    fail_name = output_name
    if fail_name.lower().endswith(".docx"):
        fail_name = f"{fail_name[:-5]}_fail.docx"
    else:
        fail_name = f"{fail_name}_fail.docx"

    try:
        default_filename = generate_default_docx(template_name, fail_name)
        return {
            "row_num": row_num,
            "output_name": output_name,
            "filename": default_filename,
            "success": False,
            "fallback": True,
            "error": last_error,
            "attempts": attempts,
        }
    except Exception as fallback_exc:
        return {
            "row_num": row_num,
            "output_name": output_name,
            "filename": None,
            "success": False,
            "fallback": False,
            "error": f"{last_error}; 回退生成失败: {str(fallback_exc)}",
            "attempts": attempts,
        }


@app.post("/generate-docx-from-excel")
async def generate_docx_from_excel(
    file: UploadFile = File(...),
    prompt: str = Form(""),
    api_url: str = Form(...),
    api_key: str = Form(...),
    model: str = Form(...),
    template_name: str = Form(...),
    sheet_name: str = Form("Sheet1"),
    prompt_column: str = Form(...),
    filename_column: str = Form(...),
    output_prefix: str = Form("excel_generated"),
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx Excel 文件")

    temp_excel_path = UPLOAD_DIR / f"excel_{uuid.uuid4().hex}.xlsx"
    with open(temp_excel_path, "wb") as f:
        f.write(await file.read())

    workbook = None
    try:
        workbook = load_workbook(filename=str(temp_excel_path), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Excel 打开失败: {str(exc)}")

    try:
        if sheet_name not in workbook.sheetnames:
            raise HTTPException(status_code=400, detail=f"未找到 sheet: {sheet_name}")

        ws = workbook[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(status_code=400, detail="Excel sheet 为空")

        prompt_col_is_letter = bool(re.fullmatch(r"[A-Za-z]+", (prompt_column or "").strip()))
        filename_col_is_letter = bool(re.fullmatch(r"[A-Za-z]+", (filename_column or "").strip()))
        use_header = not (prompt_col_is_letter and filename_col_is_letter)

        if not use_header:
            prompt_col_index = resolve_column_index(prompt_column, [])
            filename_col_index = resolve_column_index(filename_column, [])
            data_rows = list(enumerate(rows, start=1))
        else:
            header_row = [str(item or "").strip() for item in rows[0]]
            prompt_col_index = resolve_column_index(prompt_column, header_row)
            filename_col_index = resolve_column_index(filename_column, header_row)
            data_rows = list(enumerate(rows[1:], start=2))

        safe_prefix = sanitize_filename(output_prefix or "excel_generated")
        row_tasks = []

        for row_num, row in data_rows:
            prompt_text = ""
            file_name_text = ""
            if prompt_col_index < len(row):
                prompt_text = str(row[prompt_col_index] or "").strip()
            if filename_col_index < len(row):
                file_name_text = str(row[filename_col_index] or "").strip()

            if not prompt_text:
                continue

            row_prompt = build_row_prompt(prompt, prompt_text)
            if not row_prompt:
                continue

            output_name = sanitize_output_filename(file_name_text) if file_name_text else ""
            if output_name and not output_name.lower().endswith(".docx"):
                output_name = f"{output_name}.docx"
            if not output_name:
                output_name = f"{safe_prefix}_row{row_num}.docx"

            row_tasks.append({
                "row_num": row_num,
                "prompt": row_prompt,
                "output_name": output_name,
            })

        if not row_tasks:
            raise HTTPException(status_code=400, detail="提示词列没有可用数据，未生成文件")

        loop = asyncio.get_running_loop()
        with ThreadPoolExecutor(max_workers=3) as executor:
            futures = [
                loop.run_in_executor(
                    executor,
                    _process_excel_row,
                    task_info,
                    api_url,
                    api_key,
                    model,
                    template_name,
                )
                for task_info in row_tasks
            ]
            results = await asyncio.gather(*futures)

        processed_rows = len(results)
        generated_files = []
        fallback_count = 0
        success_count = 0

        summary_lines = [
            "=" * 60,
            "Excel 批量生成执行报告 (Batch Summary)",
            "=" * 60,
            f"执行时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"模板名称: {template_name}",
            f"工作表: {sheet_name}",
            f"提示词列: {prompt_column} | 文件名列: {filename_column}",
            f"并发线程数: 3",
            "-" * 60,
            f"处理总行数: {processed_rows}",
        ]

        for res in results:
            if res["filename"]:
                generated_files.append(res["filename"])
            if res["success"]:
                success_count += 1
                summary_lines.append(
                    f"[成功] 第 {res['row_num']} 行 -> {res['filename']} (尝试 {res['attempts']} 次)"
                )
            elif res["fallback"]:
                fallback_count += 1
                summary_lines.append(
                    f"[回退] 第 {res['row_num']} 行 -> {res['filename']} (原因: {res['error']})"
                )
            else:
                summary_lines.append(
                    f"[失败] 第 {res['row_num']} 行 -> 未生成文件 (原因: {res['error']})"
                )

        summary_lines.extend([
            "-" * 60,
            f"统计结果: 成功 {success_count} 个, 失败回退 {fallback_count} 个, 最终生成文件数: {len(generated_files)}",
            "=" * 60,
        ])
        summary_content = "\n".join(summary_lines)

        if not generated_files:
            raise HTTPException(status_code=400, detail="未生成任何有效文件")

        zip_filename = f"{safe_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
        zip_path = OUTPUT_DIR / zip_filename
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("batch_summary.txt", summary_content.encode("utf-8"))
            for filename in generated_files:
                source = OUTPUT_DIR / filename
                if source.exists():
                    zf.write(source, arcname=filename)

        return {
            "message": f"批量生成完成，共 {processed_rows} 个 Word，失败回退 {fallback_count} 个",
            "count": processed_rows,
            "fallback_count": fallback_count,
            "sheet_name": sheet_name,
            "prompt_column": prompt_column,
            "filename_column": filename_column,
            "zip_filename": zip_filename,
            "download_url": f"/download/{zip_filename}",
            "files": generated_files,
        }
    finally:
        if workbook is not None:
            workbook.close()
        temp_excel_path.unlink(missing_ok=True)


@app.get("/generation-records")
async def get_generation_records(limit: int = 50):
    safe_limit = max(1, min(limit, 200))
    return {"records": list_generation_records(safe_limit)}


@app.get("/generation-records/{record_id}")
async def get_generation_record_by_id(record_id: str):
    return {"record": get_generation_record(record_id)}


@app.delete("/generation-records/{record_id}")
async def delete_generation_record_endpoint(record_id: str):
    with generation_records_lock:
        if record_id not in generation_records:
            raise HTTPException(status_code=404, detail="生成记录不存在")
        del generation_records[record_id]
    dump_generation_records()
    return {"message": "记录删除成功", "record_id": record_id}


@app.delete("/generation-records")
async def clear_generation_records_endpoint():
    with generation_records_lock:
        count = len(generation_records)
        generation_records.clear()
    dump_generation_records()
    return {"message": "所有生成记录已清空", "deleted_count": count}


def _generate_and_convert_pdf(req: GenerateRequest) -> Dict[str, Any]:
    result = generate_common(req)
    docx_filename = result["docx_filename"]
    docx_path = Path(result["docx_path"])
    pdf_filename = docx_filename.replace(".docx", ".pdf")
    pdf_path = OUTPUT_DIR / pdf_filename

    convert_docx_to_pdf(docx_path, pdf_path)
    return {
        "pdf_filename": pdf_filename,
        "variables": result["variables"],
    }


@app.post("/generate-pdf")
async def generate_pdf(req: GenerateRequest):
    data = await run_in_threadpool(_generate_and_convert_pdf, req)
    pdf_filename = data["pdf_filename"]

    return {
        "message": "PDF 生成成功",
        "filename": pdf_filename,
        "download_url": f"/download/{pdf_filename}",
        "variables": data["variables"],
    }


@app.get("/download/{filename}")
async def download_file(filename: str):
    clean_filename = sanitize_filename(filename)
    file_path = OUTPUT_DIR / clean_filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="文件不存在")

    media_type = "application/octet-stream"
    if clean_filename.lower().endswith(".pdf"):
        media_type = "application/pdf"
    elif clean_filename.lower().endswith(".docx"):
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif clean_filename.lower().endswith(".zip"):
        media_type = "application/zip"

    return FileResponse(str(file_path), media_type=media_type, filename=clean_filename)


@app.exception_handler(Exception)
async def unhandled_exception_to_json(request: Request, exc: Exception):
    if isinstance(exc, HTTPException):
        return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail})
    return JSONResponse(status_code=500, content={"detail": f"Internal Server Error: {str(exc)}"})
