# -*- coding: utf-8 -*-
import main
import pytest
from docx import Document

from utils.templates import extract_invalid_template_placeholders, extract_template_variables

US_TEMPLATE = main.BASE_DIR / "US.docx"


def test_extract_template_variables_on_us_docx():
    variables = extract_template_variables(US_TEMPLATE)
    assert "story_id" in variables
    assert "story_title" in variables
    assert "story_description" in variables
    assert "story_acceptance_criteria" in variables


def test_us_docx_has_no_invalid_placeholders():
    assert extract_invalid_template_placeholders(US_TEMPLATE) == []


def test_invalid_placeholders_are_reported(tmp_path):
    doc = Document()
    doc.add_paragraph("Hello {{ valid_var }} and {{ 123invalid }} and {{ bad-identifier }}")
    target = tmp_path / "invalid.docx"
    doc.save(str(target))

    invalid = extract_invalid_template_placeholders(target)
    assert "123invalid" in invalid
    assert "bad-identifier" in invalid
    assert "valid_var" not in invalid


def test_generate_default_document_renders_blank_template(tmp_path):
    from services.document_service import generate_default_document

    generated = generate_default_document(
        template_name="US.docx",
        output_filename="blank.docx",
        upload_dir=main.UPLOAD_DIR,
        output_dir=tmp_path,
    )
    assert generated == "blank.docx"
    assert (tmp_path / "blank.docx").exists()


if __name__ == "__main__":
    pytest.main([__file__])
