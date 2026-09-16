import re
from pathlib import Path
from typing import List

from docx import Document
from docxtpl import DocxTemplate


def extract_template_text(template_path: Path) -> str:
    doc = Document(str(template_path))
    text_blocks = []

    for paragraph in doc.paragraphs:
        text_blocks.append(paragraph.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text_blocks.append(paragraph.text)

    return "\n".join(text_blocks)


def extract_template_variables_regex(template_path: Path) -> List[str]:
    full_text = extract_template_text(template_path)
    vars_found = re.findall(r"{{\s*([a-zA-Z_][a-zA-Z0-9_]*)\s*}}", full_text)

    ordered: List[str] = []
    for key in vars_found:
        if key not in ordered:
            ordered.append(key)

    return ordered


def extract_template_variables(template_path: Path) -> List[str]:
    try:
        tpl = DocxTemplate(str(template_path))
        undeclared = tpl.get_undeclared_template_variables()
        if undeclared:
            full_text = extract_template_text(template_path)
            ordered: List[str] = []
            for token in re.findall(r"[a-zA-Z_][a-zA-Z0-9_]*", full_text):
                if token in undeclared and token not in ordered:
                    ordered.append(token)
            for token in sorted(undeclared):
                if token not in ordered:
                    ordered.append(token)
            return ordered
        return extract_template_variables_regex(template_path)
    except Exception:
        return extract_template_variables_regex(template_path)

def extract_invalid_template_placeholders(template_path: Path) -> List[str]:
    full_text = extract_template_text(template_path)
    all_placeholders = re.findall(r"{{\s*([^{}]+?)\s*}}", full_text)

    invalid = []
    for placeholder in all_placeholders:
        if not re.fullmatch(r"[a-zA-Z_][a-zA-Z0-9_]*", placeholder):
            if placeholder not in invalid:
                invalid.append(placeholder)

    return invalid
